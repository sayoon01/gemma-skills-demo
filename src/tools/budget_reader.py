"""예실대비표를 읽고 원본 셀 위치와 검증 결과를 JSON으로 저장한다."""

import argparse
import hashlib
import json
import re
import tempfile
import warnings
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]


def normalize_label(value):
    """비교용 헤더에서 공백을 제거한다."""
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value))


def parse_number(value):
    """숫자와 숫자 문자열을 Decimal로 변환한다. 빈 값은 보존한다."""
    if value is None:
        return None

    if isinstance(value, bool):
        raise ValueError("논리값은 금액으로 변환하지 않습니다.")

    text = str(value).strip()
    if not text:
        return None

    # 쉼표가 있으면 올바른 천 단위 구분인지 확인한다.
    if "," in text:
        pattern = r"[+-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?"
        if not re.fullmatch(pattern, text):
            raise ValueError(f"숫자 구분 형식이 잘못됐습니다: {value!r}")
        text = text.replace(",", "")

    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(
            f"숫자로 해석할 수 없는 값: {value!r}"
        ) from exc

    if not number.is_finite():
        raise ValueError(f"유한하지 않은 숫자: {value!r}")

    return number


def json_number(value):
    """정수는 JSON 숫자, 소수는 정밀도를 보존하는 문자열로 기록한다."""
    if value is None:
        return None

    if value == value.to_integral_value():
        return int(value)

    return str(value)


def normalize_code(value):
    """문자형 코드의 앞자리 0은 보존한다."""
    if isinstance(value, bool):
        raise ValueError("논리값은 비용코드로 사용할 수 없습니다.")

    if isinstance(value, (int, float)):
        number = parse_number(value)
        if number != number.to_integral_value():
            raise ValueError(f"정수가 아닌 비용코드: {value!r}")
        return str(int(number))

    text = str(value).strip()
    if not text:
        raise ValueError("비용코드가 비어 있습니다.")

    return text


def merged_value(sheet, row, column):
    """병합 셀은 왼쪽 위 기준 셀의 값을 읽는다."""
    cell = sheet.cell(row, column)

    if cell.value is not None:
        return cell.value

    for region in sheet.merged_cells.ranges:
        if (
            region.min_row <= row <= region.max_row
            and region.min_col <= column <= region.max_col
        ):
            return sheet.cell(
                region.min_row,
                region.min_col,
            ).value

    return None


def locate_headers(sheet):
    """상단의 예실대비표 헤더를 찾아 2단 헤더를 결합한다."""
    header_row = None

    for row in range(1, min(sheet.max_row, 20) + 1):
        labels = {
            normalize_label(cell.value)
            for cell in sheet[row]
            if cell.value is not None
        }

        if {"비목분류", "비용명", "계획예산"}.issubset(labels):
            header_row = row
            break

    if header_row is None:
        raise ValueError("예실대비표 헤더를 찾지 못했습니다.")

    columns = {}

    for column in range(1, sheet.max_column + 1):
        top = normalize_label(
            merged_value(sheet, header_row, column)
        )
        bottom = normalize_label(
            merged_value(sheet, header_row + 1, column)
        )

        if bottom and bottom != top:
            label = f"{top}/{bottom}"
        else:
            label = top

        if label:
            columns.setdefault(label, []).append(column)

    def unique(label):
        candidates = columns.get(label, [])

        if len(candidates) != 1:
            raise ValueError(
                f"헤더 {label!r}를 하나로 식별할 수 없습니다: "
                f"{candidates}"
            )

        return candidates[0]

    # 해당 예실대비표 형식의 비용명 헤더는 코드·이름 두 열을 포함한다.
    cost_columns = columns.get("비용명", [])

    if len(cost_columns) != 2:
        raise ValueError(
            "비용명 헤더 아래의 코드·이름 열 구조가 예상과 다릅니다."
        )

    mapping = {
        "category": unique("비목분류"),
        "code": cost_columns[0],
        "name": cost_columns[1],
    }

    numeric_columns = {}

    for label, candidates in columns.items():
        if label in {"비목분류", "비용명"}:
            continue

        if len(candidates) != 1:
            raise ValueError(f"중복 숫자 헤더: {label}")

        numeric_columns[label] = candidates[0]

    required_metrics = [
        "실행예산/합계",
        "집행계/합계",
        "예산잔액/합계",
    ]

    for label in required_metrics:
        if label not in numeric_columns:
            raise ValueError(f"필수 지표 열이 없습니다: {label}")

    return header_row + 2, mapping, numeric_columns


def inspect_sheet(sheet):
    start_row, mapping, numeric_columns = locate_headers(sheet)

    details = []
    summaries = []
    issues = []
    converted_cells = []
    current_category = None
    seen_codes = set()

    for row in range(start_row, sheet.max_row + 1):
        if not any(
            cell.value is not None
            for cell in sheet[row]
        ):
            continue

        category_value = sheet.cell(
            row, mapping["category"]
        ).value
        category_label = normalize_label(category_value)

        code_value = sheet.cell(row, mapping["code"]).value
        name_value = sheet.cell(row, mapping["name"]).value

        if category_label in {
            "소계",
            "내부흡수액",
            "외부유출액",
            "합계",
        }:
            row_type = category_label

        elif (
            code_value is not None
            and name_value is not None
            and str(code_value).strip()
            and str(name_value).strip()
        ):
            row_type = "detail"

        else:
            issues.append({
                "row": row,
                "message": "행 유형을 식별할 수 없습니다.",
            })
            continue

        if row_type == "detail":
            if category_label:
                current_category = str(category_value).strip()

            if current_category is None:
                issues.append({
                    "row": row,
                    "message": "상세행의 비목분류를 확인할 수 없습니다.",
                })

        values = {}
        locations = {}

        for label, column in numeric_columns.items():
            cell = sheet.cell(row, column)
            location = f"{sheet.title}!{cell.coordinate}"
            locations[label] = location

            if cell.data_type == "f":
                values[label] = None
                issues.append({
                    "cell": location,
                    "message": (
                        "수식 셀입니다. 이번 읽기 단계에서는 "
                        "수식 평가나 캐시값 사용을 하지 않습니다."
                    ),
                })
                continue

            try:
                number = parse_number(cell.value)
                values[label] = json_number(number)

                if isinstance(cell.value, str) and number is not None:
                    converted_cells.append(location)

            except ValueError as exc:
                values[label] = None
                issues.append({
                    "cell": location,
                    "message": str(exc),
                })

        item = {
            "row": row,
            "row_type": row_type,
            "values": values,
            "source_cells": locations,
        }

        if row_type == "detail":
            try:
                code = normalize_code(code_value)
            except ValueError as exc:
                code = str(code_value)
                issues.append({
                    "row": row,
                    "message": str(exc),
                })

            item.update({
                "category": current_category,
                "code": code,
                "name": str(name_value).strip(),
            })

            if code in seen_codes:
                issues.append({
                    "row": row,
                    "message": f"비용코드 중복: {code}",
                })

            seen_codes.add(code)
            details.append(item)

        else:
            item["label"] = str(category_value)
            summaries.append(item)

            if row_type == "소계":
                current_category = None

    # 상세행만 합산하여 소계·요약행의 중복 집계를 피한다.
    total_rows = [
        item for item in summaries
        if item["row_type"] == "합계"
    ]
    reconciliation = []

    if not details:
        issues.append({
            "message": "상세행을 찾지 못했습니다.",
        })

    if len(total_rows) != 1:
        issues.append({
            "message": "전체 합계행을 하나로 식별하지 못했습니다.",
        })

    elif details:
        total = total_rows[0]

        for label in numeric_columns:
            operands = [
                item["values"][label]
                for item in details
            ]
            original = total["values"][label]

            if original is None or any(
                value is None for value in operands
            ):
                reconciliation.append({
                    "metric": label,
                    "status": "not_checked",
                    "reason": "비어 있거나 해석하지 못한 값이 있습니다.",
                })
                continue

            calculated = sum(
                (Decimal(str(value)) for value in operands),
                Decimal("0"),
            )
            expected = Decimal(str(original))

            reconciliation.append({
                "metric": label,
                "detail_sum": json_number(calculated),
                "reported_total": original,
                "difference": json_number(calculated - expected),
                "matched": calculated == expected,
                "total_source": total["source_cells"][label],
            })

    negative_balances = []

    for item in details:
        for metric, value in item["values"].items():
            if (
                metric.startswith("예산잔액/")
                and value is not None
                and Decimal(str(value)) < 0
            ):
                negative_balances.append({
                    "code": item["code"],
                    "name": item["name"],
                    "metric": metric,
                    "value": value,
                    "source": item["source_cells"][metric],
                    "interpretation": (
                        "음수 잔액을 확인했습니다. "
                        "업무상 오류 여부는 별도 판단이 필요합니다."
                    ),
                })

    return {
        "sheet": sheet.title,
        "dimensions": {
            "rows": sheet.max_row,
            "columns": sheet.max_column,
        },
        "numeric_headers": list(numeric_columns),
        "detail_count": len(details),
        "details": details,
        "summary_rows": summaries,
        "numeric_text_conversions": converted_cells,
        "total_reconciliation": reconciliation,
        "negative_balances": negative_balances,
        "issues": issues,
    }


def inspect_workbook(path):
    path = Path(path).resolve()
    original_hash = hashlib.sha256(path.read_bytes()).hexdigest()

    report = {
        "file": path.name,
        "sha256": original_hash,
        "sheets": [],
    }

    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")

        workbook = load_workbook(path, data_only=False)

        try:
            report["external_link_count"] = len(
                workbook._external_links
            )
            report["formula_count"] = sum(
                cell.data_type == "f"
                for sheet in workbook
                for row in sheet
                for cell in row
            )

            for sheet in workbook:
                report["sheets"].append(
                    inspect_sheet(sheet)
                )

        finally:
            workbook.close()

    report["reader_warnings"] = sorted({
        str(warning.message)
        for warning in caught
    })

    current_hash = hashlib.sha256(path.read_bytes()).hexdigest()

    if current_hash != original_hash:
        raise RuntimeError("검사 중 원본 파일 내용이 변경됐습니다.")

    report["source_hash_unchanged"] = True
    return report


def main():
    parser = argparse.ArgumentParser(
        description="예실대비표 읽기 및 상세합·전체합 대조"
    )
    parser.add_argument(
        "files",
        nargs="+",
        type=Path,
        help="검사할 예실대비표 파일 경로",
    )
    args = parser.parse_args()

    output_root = ROOT / "outputs" / "budget-inspection"
    output_root.mkdir(parents=True, exist_ok=True)

    run_dir = Path(
        tempfile.mkdtemp(prefix="run-", dir=output_root)
    )

    reports = []
    needs_review = False

    for path in args.files:
        try:
            report = inspect_workbook(path)
            reports.append(report)

            print(f"\n파일: {report['file']}")
            print(f"수식 개수: {report['formula_count']}")
            print("원본 변경 없음: 확인")

            for warning in report["reader_warnings"]:
                print(f"읽기 경고: {warning}")

            for sheet in report["sheets"]:
                mismatches = [
                    item
                    for item in sheet["total_reconciliation"]
                    if item.get("matched") is False
                ]
                unchecked = [
                    item
                    for item in sheet["total_reconciliation"]
                    if item.get("status") == "not_checked"
                ]

                sheet_needs_review = bool(
                    sheet["issues"]
                    or mismatches
                    or unchecked
                )
                needs_review = needs_review or sheet_needs_review

                print(f"시트: {sheet['sheet']}")
                print(f"상세행: {sheet['detail_count']}개")
                print(f"요약행: {len(sheet['summary_rows'])}개")
                print(
                    "숫자 문자열 변환:",
                    len(sheet["numeric_text_conversions"]),
                    "셀",
                )
                print("상세합·전체합 불일치:", len(mismatches))
                print("대조 미실시:", len(unchecked))
                print(
                    "음수 잔액:",
                    len(sheet["negative_balances"]),
                )
                print("읽기 확인 사항:", len(sheet["issues"]))

                for item in mismatches:
                    print(
                        f"  합계 확인 필요: {item['metric']} / "
                        f"차이={item['difference']}"
                    )

                for issue in sheet["issues"]:
                    print(
                        "  확인 필요:",
                        json.dumps(issue, ensure_ascii=False),
                    )

        except Exception as exc:
            needs_review = True
            reports.append({
                "file": path.name,
                "error": f"{type(exc).__name__}: {exc}",
            })
            print(f"\n실패: {path.name} / {exc}")

    output = {
        "checked_at": datetime.now(timezone.utc).isoformat(),
        "status": "needs_review" if needs_review else "passed",
        "scope": "원본 읽기·상세행 추출·전체 합계 대조",
        "model_used": False,
        "source_write_performed": False,
        "limitations": [
            "각 소계행의 집계 정확성은 아직 검사하지 않습니다.",
            "내부흡수액과 외부유출액의 분류 기준은 검증하지 않습니다.",
            "음수 잔액을 업무상 오류로 단정하지 않습니다.",
            "수식 셀을 평가하거나 재계산하지 않습니다.",
            "다른 양식의 엑셀을 자동으로 해석하는 범용 도구는 아닙니다.",
        ],
        "workbooks": reports,
    }

    output_file = run_dir / "inspection.json"
    output_file.write_text(
        json.dumps(output, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    print(f"\n최종 상태: {output['status']}")
    print(f"결과 파일: {output_file}")

    return 1 if needs_review else 0


if __name__ == "__main__":
    raise SystemExit(main())
